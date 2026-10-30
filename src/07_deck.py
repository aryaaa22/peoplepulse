"""Phase 4 deliverable: HR planning recommendation deck (PPTX) built from the processed outputs."""
import json

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from config import CHARTS, OUT, PROC

NAVY, CORAL, TEAL, GREY, LIGHT, WHITE, AMBER = (RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0xE7, 0x6F, 0x51), RGBColor(0x2A, 0x9D, 0x8F),
                                                RGBColor(0x5C, 0x67, 0x7D), RGBColor(0xF3, 0xF5, 0xF8), RGBColor(0xFF, 0xFF, 0xFF),
                                                RGBColor(0xB4, 0x53, 0x09))
W, H = Inches(13.333), Inches(7.5)


def load():
    m = json.loads((PROC / "model_metrics.json").read_text())
    d = pd.read_csv(PROC / "attrition_diagnostics.csv")
    sc = pd.read_csv(PROC / "talent_health_scorecard.csv")
    cr = pd.read_csv(PROC / "critical_role_risk.csv")
    fc = pd.read_csv(PROC / "headcount_forecast.csv")
    fl = pd.read_csv(PROC / "early_warning_flag_performance.csv")
    rs = pd.read_csv(PROC / "risk_scores.csv")
    return m, d, sc, cr, fc, fl, rs


def text(slide, x, y, w, h, s, size=14, bold=False, color=NAVY, bullets=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    items = bullets if bullets is not None else [s]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ("•  " if bullets is not None else "") + item
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"
        p.space_after = Pt(8 if bullets is not None else 0)
    return tb


def base(prs, title, kicker=None, synthetic=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = AMBER if synthetic else NAVY; bar.line.fill.background()
    if kicker:
        text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4), kicker.upper(), 11, True, AMBER if synthetic else TEAL)
    text(s, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.9), title, 26, True)
    text(s, Inches(0.6), Inches(7.05), Inches(12), Inches(0.3),
         "PeoplePulse | IBM HR Analytics (1,470 records)" + (" | SYNTHETIC DATA – illustrative only" if synthetic else ""), 9, color=GREY)
    return s


def stat(slide, x, y, value, label, color=NAVY):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.85), Inches(1.35))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.fill.background()
    box.adjustments[0] = 0.08
    text(slide, x + Inches(0.2), y + Inches(0.12), Inches(2.5), Inches(0.6), value, 30, True, color)
    text(slide, x + Inches(0.2), y + Inches(0.78), Inches(2.5), Inches(0.5), label, 12, color=GREY)


def picture(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(CHARTS / f"{name}.png"), x, y, width=w, height=h)


def table(slide, df, x, y, w, col_w=None, font=11, row_h=0.36):
    rows, cols = df.shape[0] + 1, df.shape[1]
    t = slide.shapes.add_table(rows, cols, x, y, w, Inches(row_h * rows)).table
    for j, c in enumerate(df.columns):
        t.cell(0, j).text = str(c)
    for i, r in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(r):
            t.cell(i, j).text = str(v)
    for i in range(rows):
        for j in range(cols):
            cell = t.cell(i, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font); run.font.name = "Calibri"; run.font.bold = i == 0
                    run.font.color.rgb = WHITE if i == 0 else NAVY
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    return t


def main():
    m, d, sc, cr, fc, fl, rs = load()
    n, exits = int(d.query("dimension == 'Department'").records.sum()), int(d.query("dimension == 'Department'").exits.sum())
    active = n - exits
    ann = exits / ((n + active) / 2)
    get = lambda dim, val: d[(d.dimension == dim) & (d.value.astype(str) == val)].iloc[0]
    ot, ot_no = get("Overtime", "Yes"), get("Overtime", "No")
    t01, newm, srep = get("Tenure band", "0–1 Years"), get("Manager tenure", "New manager (<1 yr)"), get("Job role", "Sales Representative")
    lab = get("Job role", "Laboratory Technician")
    lr = m["logistic_regression"]
    hi = next(t for t in m["tier_performance"] if t["tier"] == "High")
    allfc = fc[fc.department == "All"]
    last = allfc.groupby("scenario").last()
    exits4 = allfc.groupby("scenario").expected_exits.sum()
    base_s, ret_s, frz_s = "Base: backfill 80% of exits", "Retention programme (-30% risk for High tier) + 80% backfill", "Hiring freeze"
    saved = exits4[base_s] - exits4[ret_s]
    se = cr[cr.job_role == "Sales Executive"].iloc[0]
    sales = sc[sc.department == "Sales"].iloc[0]
    or_ot = pd.read_csv(PROC / "risk_drivers.csv").set_index("feature").odds_ratio["overtime_Yes"]
    high_active = int(round((sc.high_risk_share * sc.active).sum()))
    f4 = fl[fl.flag_bucket.astype(str) == "4+"].iloc[0]; f0 = fl[fl.flag_bucket.astype(str) == "0"].iloc[0]

    prs = Presentation(); prs.slide_width, prs.slide_height = W, H

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    text(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(1.2), "PeoplePulse", 54, True, WHITE)
    text(s, Inches(0.8), Inches(3.4), Inches(11.5), Inches(0.8), "Workforce attrition & productivity analytics – HR planning recommendations", 24, False, WHITE)
    text(s, Inches(0.8), Inches(4.5), Inches(11.5), Inches(0.5), "Data Analytics internship project | IBM HR Analytics dataset (1,470 employees) + labelled synthetic operational layer", 14, False, RGBColor(0xC9, 0xD3, 0xE0))

    # 2 Executive summary
    s = base(prs, "Attrition is concentrated, predictable and largely addressable", "Executive summary")
    stat(s, Inches(0.6), Inches(1.6), f"{ann:.1%}", f"annualised attrition ({exits} exits / {n:,} records)", CORAL)
    stat(s, Inches(3.65), Inches(1.6), f"{ot.exit_proportion:.0%}", f"exit rate with overtime vs {ot_no.exit_proportion:.0%} without", CORAL)
    stat(s, Inches(6.7), Inches(1.6), f"{t01.exit_proportion:.0%}", "exit rate at 0–1 years' tenure", CORAL)
    stat(s, Inches(9.75), Inches(1.6), f"{hi['share_of_all_exits']:.0%}", "of exits fall in the model's top-10% risk tier", TEAL)
    text(s, Inches(0.6), Inches(3.3), Inches(12.2), Inches(3.6), "", 15, bullets=[
        f"Overtime is the single strongest driver (odds ratio ×{or_ot:.1f} after controlling for role, pay, tenure and satisfaction).",
        f"Risk peaks early: 0–1 yr tenure ({t01.exit_proportion:.0%}), new manager in the last year ({newm.exit_proportion:.0%}), grade G1 and age < 25.",
        f"Sales Representatives ({srep.exit_proportion:.0%} exit) and Laboratory Technicians ({lab.exit_proportion:.0%}) are the hotspot roles; Sales scores Red on talent health ({sales.talent_health_score:.0f}/100).",
        f"Without action, active headcount drifts from {active:,} to ~{last.loc[base_s, 'headcount_p50']:.0f} by 2027Q3 even with 80% backfill; a freeze would take it to ~{last.loc[frz_s, 'headcount_p50']:.0f}.",
        f"Recommended: overtime controls, a first-year / manager-transition programme, pay review for entry roles and a targeted retention plan for {high_active} high-risk employees.",
    ])

    # 3 Data foundation
    s = base(prs, "Data foundation: what is real, what is synthetic, what was fixed", "Phase 1 – HR data foundation")
    text(s, Inches(0.6), Inches(1.6), Inches(6), Inches(0.4), "Real – IBM HR Analytics (1,470 records)", 16, True, TEAL)
    text(s, Inches(0.6), Inches(2.05), Inches(6), Inches(2.4), "", 13, bullets=[
        "Employee master: department, role, grade, tenure, pay, overtime, satisfaction, rating, manager & promotion history",
        "Drives all attrition, risk-model, scorecard, critical-role and forecast results",
        "No hire/exit dates or locations – left blank, not invented; rates assume exits span 12 months",
    ])
    text(s, Inches(6.9), Inches(1.6), Inches(6), Inches(0.4), "Synthetic – generated, flagged is_synthetic", 16, True, AMBER)
    text(s, Inches(6.9), Inches(2.05), Inches(6), Inches(2.4), "", 13, bullets=[
        "12 months of attendance, 4 quarters of productivity scores, 215 requisitions with source & time-to-fill",
        "Driven only by real IBM attributes; no exit signal injected",
        "Used to demonstrate absenteeism, hiring-source and productivity analytics – not as evidence",
    ])
    table(s, pd.DataFrame([
        ("KPI_Summary", "'Exited employee records' used COUNTA → 1,470", "COUNTIF(Attrition=\"Yes\") → 237"),
        ("Attendance / Performance / Hiring samples", "IDs E0003/6/9/17 not in master; 5 hires collide with long-tenured IBM IDs", "Replaced by synthetic layer keyed to real IDs"),
        ("Dimensions", "Tenure bands and grade list (G1–G4) don't match data (G5 present)", "Dimensions rebuilt from data"),
    ], columns=["Source sheet", "Issue found", "Resolution"]), Inches(0.6), Inches(4.55), Inches(12.1), [3.2, 5.2, 3.7], 11)

    # 4 Hotspots
    s = base(prs, "Who leaves: early-career, overtime-heavy, entry-grade staff", "Phase 2 – attrition & risk analytics")
    picture(s, "attrition_hotspots", Inches(0.5), Inches(1.55), w=Inches(7.6))
    text(s, Inches(8.4), Inches(1.7), Inches(4.5), Inches(5), "", 13, bullets=[
        "Only cohorts whose 95% confidence interval sits above the company rate are shown.",
        f"Sales Representatives: {srep.exit_proportion:.0%} exit proportion (n={int(srep.records)}).",
        f"Overtime workers: {ot.exit_proportion:.0%} vs {ot_no.exit_proportion:.0%}; overtime + low satisfaction reaches 37%.",
        f"New manager in the past year: {newm.exit_proportion:.0%} – a proxy for manager change (no dated manager history).",
        "Seniority protects: G4–G5, Managers and Directors exit at 2–7%.",
    ])

    # 5 Cohort
    s = base(prs, "Risk is front-loaded: the first two years drive exits in every department", "Phase 2 – cohort analysis")
    picture(s, "cohort_heatmap", Inches(0.5), Inches(1.6), w=Inches(8.2))
    text(s, Inches(8.9), Inches(1.7), Inches(4.1), Inches(5), "", 13, bullets=[
        "R&D: 32% exit in 0–1 yr vs 5% at 10+ yrs.",
        "Sales stays elevated through 2–3 yrs (29%) – the onboarding problem extends into year 3.",
        "HR 0–1 yr = 80% but n = 5: directional only.",
        "Implication: retention spend has the highest return in years 0–3.",
    ])

    # 6 Early warning
    s = base(prs, "An early-warning score flags 44% of leavers in 10% of the workforce", "Phase 2 – early-warning indicators")
    picture(s, "risk_drivers", Inches(0.4), Inches(1.5), w=Inches(6.3))
    picture(s, "early_warning_flags", Inches(6.9), Inches(1.5), w=Inches(5.9))
    text(s, Inches(6.9), Inches(5.15), Inches(6), Inches(1.8), "", 12, bullets=[
        f"Logistic regression, 5-fold CV: ROC AUC {lr['roc_auc']}, PR AUC {lr['pr_auc']} vs {m['overall_exit_proportion']:.0%} baseline; gradient boosting no better ({m['gradient_boosting']['roc_auc']}).",
        f"Simple rule version: employees with 4+ of 7 flags exit at {f4.exit_proportion:.0%} vs {f0.exit_proportion:.0%} with none – usable without the model.",
    ])

    # 7 Critical roles
    s = base(prs, "Critical-role exposure: Sales Executives are the main continuity risk", "Phase 4 – critical role risk")
    picture(s, "critical_roles", Inches(0.5), Inches(1.55), w=Inches(7.4))
    text(s, Inches(8.2), Inches(1.7), Inches(4.8), Inches(5), "", 13, bullets=[
        f"Sales Executive: {int(se.active)} active, {int(se.high_risk_active)} high-risk, ~{se.expected_exits_12m:.0f} expected exits in 12 months.",
        "Leadership roles (Manager, Research / Manufacturing Director) are stable: ≤ 9 expected exits each.",
        "Lab Technicians and Research Scientists are not tagged critical but together lose ~70 people a year – a capacity risk for R&D.",
        "Critical = leadership or revenue-carrying roles (assumption; adjust to company policy).",
    ])

    # 8 Scorecard
    s = base(prs, "Department talent health scorecard", "Phase 4 – talent health")
    df = sc.assign(**{
        "Department": sc.department, "Score": sc.talent_health_score.round(0).astype(int).astype(str) + " " + sc.rag,
        "Attrition": (sc.attrition_rate * 100).round(1).astype(str) + "%", "High-risk share": (sc.high_risk_share * 100).round(1).astype(str) + "%",
        "Low satisfaction": (sc.low_satisfaction_share * 100).round(1).astype(str) + "%", "Overtime": (sc.overtime_share * 100).round(1).astype(str) + "%",
        "Promotion stalled 5+ yrs": (sc.promotion_stalled_share * 100).round(1).astype(str) + "%",
        "Absenteeism*": (sc.absenteeism_rate_syn * 100).round(1).astype(str) + "%", "Time to fill*": sc.avg_time_to_fill_syn.round(0).astype(int).astype(str) + " d",
    })[["Department", "Score", "Attrition", "High-risk share", "Low satisfaction", "Overtime", "Promotion stalled 5+ yrs", "Absenteeism*", "Time to fill*"]]
    table(s, df, Inches(0.6), Inches(1.7), Inches(12.1), [2.0, 1.1, 1.0, 1.2, 1.4, 1.2, 1.6, 1.5, 1.1], 12, 0.5)
    text(s, Inches(0.6), Inches(4.1), Inches(12.1), Inches(2.6), "", 13, bullets=[
        "Score 0–100: each metric scored vs the company level (100 at or better, −2 pts per 1% worse); weights: attrition 30%, high-risk 20%, low satisfaction 15%, overtime 10%, absenteeism 10%, time to fill 10%, promotion 5%.",
        "Sales is Red: highest attrition (23%) and 2.3× R&D's high-risk share. HR is Amber on low satisfaction (47%) and attrition.",
        "* synthetic inputs (10% + 10% of weight).",
    ])

    # 9 Forecast
    s = base(prs, "Quarterly headcount forecast: backfilling alone does not hold headcount", "Phase 4 – workforce planning")
    picture(s, "headcount_forecast", Inches(0.4), Inches(1.5), w=Inches(7.6))
    tdf = pd.DataFrame({"Scenario": [x.split(" (")[0].split(":")[0] for x in last.index], "2027Q3 median": last.headcount_p50.round(0).astype(int).values,
                        "P10–P90": [f"{a:.0f}–{b:.0f}" for a, b in zip(last.headcount_p10, last.headcount_p90)], "Exits (4 qtrs)": exits4[last.index].round(0).astype(int).values})
    table(s, tdf, Inches(8.2), Inches(1.7), Inches(4.8), [1.8, 1.0, 1.1, 0.9], 11)
    text(s, Inches(8.2), Inches(3.9), Inches(4.8), Inches(3), "", 12, bullets=[
        f"Starting point: {active:,} active employees (2026Q3 snapshot).",
        f"Retention programme on the High tier avoids ~{saved:.0f} exits in a year.",
        "Holding headcount flat needs ~45 hires per quarter; growing 2% per quarter needs ~70.",
        "Method: 2,000 Monte Carlo runs over each employee's risk score; backfills arrive a quarter after the exit; new hires carry the observed first-year exit rate.",
    ])

    # 10 Synthetic
    s = base(prs, "Hiring, absence and productivity analytics (method demo)", "Phase 3 – productivity & hiring insights", synthetic=True)
    picture(s, "hiring_source_syn", Inches(0.4), Inches(1.5), w=Inches(6.2))
    picture(s, "absenteeism_syn", Inches(6.8), Inches(1.5), w=Inches(6.1))
    text(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(2), "", 12, bullets=[
        "Synthetic: the IBM dataset has no attendance, requisition or productivity history. Values show what the dashboard and KPIs will report once real HRIS / ATS data is connected.",
        "Real component: first-year exits (16 of 44 employees with < 1 yr tenure = 36%) – the early-exit outcome attached to each requisition.",
        "Ready-built: absenteeism by month & department, time-to-fill by source & grade, probation outcomes, productivity by role family and rating.",
    ], color=NAVY)

    # 11 Recommendations
    s = base(prs, "Recommendations for HR and leadership", "HR planning recommendations")
    rec = pd.DataFrame([
        ("1. Cap and rebalance overtime", f"{ot.exit_proportion:.0%} vs {ot_no.exit_proportion:.0%} exit; OR ×{or_ot:.1f}; 76% of High-tier actives work overtime", "Overtime caps / comp-time for Sales Reps, Lab Techs, Research Scientists; monthly overtime review", "Ops + HRBP"),
        ("2. First-two-years programme", f"0–1 yr exit {t01.exit_proportion:.0%}; Sales 2–3 yr 29%", "Structured onboarding, 30/60/90-day check-ins, buddy system, 12-month career conversation", "HR + line managers"),
        ("3. Protect manager transitions", f"New manager in last year: {newm.exit_proportion:.0%} exit", "Handover plan and skip-level 1:1 within 30 days of any manager change", "HRBP"),
        ("4. Entry-role pay & progression review", "G1 exit 26%; leavers in Sales Rep / Lab Tech earn 12–15% less than stayers", "Benchmark G1 pay; clearer G1→G2 promotion criteria", "Total Rewards"),
        ("5. Targeted retention for High tier", f"{high_active} high-risk active staff; ~{saved:.0f} exits avoidable / yr", "Monthly risk-score refresh; stay interviews; manager action plans", "HR Analytics + HRBP"),
        ("6. Sales Executive bench", f"~{se.expected_exits_12m:.0f} expected exits / yr in a critical role", "Succession pool from top Sales Reps; pipeline hiring ahead of exits", "Sales leadership + TA"),
    ], columns=["Action", "Evidence", "What to do", "Owner"])
    table(s, rec, Inches(0.5), Inches(1.55), Inches(12.3), [2.6, 3.6, 4.4, 1.7], 11, 0.75)

    # 12 Next steps
    s = base(prs, "Next steps and data needed to go from snapshot to live system", "Roadmap")
    text(s, Inches(0.6), Inches(1.6), Inches(6), Inches(5), "", 14, bullets=[
        "Connect HRIS: hire / exit dates, location, monthly snapshots → true period attrition and trend lines",
        "Replace synthetic layer with real attendance (time & attendance system) and requisitions (ATS)",
        "Define productivity per role family (e.g. quota attainment for Sales, output KPIs for R&D)",
        "Re-train the risk model quarterly; validate on the next period's actual exits",
    ])
    text(s, Inches(6.9), Inches(1.6), Inches(6), Inches(5), "", 14, bullets=[
        "Week 1–2: pilot retention actions with Sales and the R&D lab teams",
        "Track KPIs monthly: attrition, first-year exit rate, overtime share, High-tier count",
        "Governance: risk scores are for workforce planning in aggregate – not for individual performance decisions",
        "Fairness check: gender is a model input – audit tier rates by gender / age and test removing it before operational use",
    ])

    # 13 Appendix
    s = base(prs, "Appendix: methodology and KPI definitions", "Appendix")
    table(s, pd.DataFrame([
        ("Annualised attrition rate", "Exits / average headcount = exits / ((records + active) / 2); assumes IBM exits span 12 months", "Real"),
        ("Exit proportion", "Exits / employee records in a segment; 95% Wilson confidence interval", "Real"),
        ("Risk probability", "Out-of-fold logistic regression probability (10 numeric + 4 categorical features)", "Real"),
        ("Risk tier", f"High ≥ {m['tier_thresholds']['high']:.2f} (top 10% of active), Medium ≥ {m['tier_thresholds']['medium']:.2f} (next 20%)", "Real"),
        ("Absenteeism rate", "Absent days / scheduled workdays", "Synthetic"),
        ("Time to fill", "Fill date − opening date (days)", "Synthetic"),
        ("Early exit within 1 year", "Exited with Years_at_Company = 0", "Real"),
        ("Talent health score", "Weighted 0–100 composite vs company level (see scorecard slide)", "Mixed"),
    ], columns=["KPI", "Definition", "Data basis"]), Inches(0.6), Inches(1.6), Inches(12.1), [2.8, 7.9, 1.4], 11, 0.5)

    path = OUT / "PeoplePulse_HR_Recommendations.pptx"
    prs.save(path)
    print("wrote", path, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
