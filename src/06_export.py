"""Export: Power BI / Excel-ready workbook, dashboard JSON and static charts for the deck."""
import json

import duckdb
import matplotlib
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

from config import CHARTS, DB, OUT, PROC, ROOT, SNAPSHOT_DATE, STAGE  # noqa: E402

NAVY, TEAL, CORAL, GREY, AMBER = "#1F3A5F", "#2A9D8F", "#E76F51", "#8D99AE", "#E9C46A"


def q(con, sql):
    return con.execute(sql).df()


def workbook(con):
    dq = pd.DataFrame(json.loads((STAGE / "data_quality.json").read_text()))
    k = q(con, "SELECT * FROM vw_kpi_overall").iloc[0]
    master = q(con, """SELECT m.*, r.risk_probability, r.risk_tier, r.top_risk_drivers, f.flag_count AS early_warning_flag_count
                       FROM stg_employee_master m JOIN an_risk_scores r USING (employee_id)
                       JOIN an_early_warning_flags f USING (employee_id) ORDER BY employee_id""")
    em_cols = {c: get_column_letter(i + 1) for i, c in enumerate(master.columns)}
    n = len(master) + 1
    kpi = pd.DataFrame([
        ("Employee records", f"=COUNTA(Employee_Master!{em_cols['employee_id']}2:{em_cols['employee_id']}{n})", "Real (IBM)", "Count of employee master records"),
        ("Active headcount", f'=COUNTIF(Employee_Master!{em_cols["employment_status"]}2:{em_cols["employment_status"]}{n},"Active")', "Real (IBM)", "Employment_Status = Active"),
        ("Exited employee records", f'=COUNTIF(Employee_Master!{em_cols["attrition"]}2:{em_cols["attrition"]}{n},"Yes")', "Real (IBM)", "Attrition = Yes (fixes COUNTA bug in source workbook)"),
        ("Dataset attrition proportion", "=B4/B2", "Real (IBM)", "Exits / records"),
        ("Annualised attrition rate", "=B4/((B2+B3)/2)", "Real (IBM) + assumption", "Exits / average headcount; assumes IBM exits span 12 months"),
        ("Absenteeism rate", float(k.absenteeism_rate_syn), "SYNTHETIC", "Absent days / scheduled workdays, last 12 months"),
        ("Average time to fill (days)", float(k.avg_time_to_fill_days_syn), "SYNTHETIC", "Fill date - opening date, recent hires"),
        ("Average productivity score (latest quarter)", float(k.avg_productivity_latest_q_syn), "SYNTHETIC", "0–100 score"),
    ], columns=["KPI", "Value", "Data basis", "Definition"])
    dims = {
        "Dim_Department": q(con, "SELECT * FROM dim_department"),
        "Dim_Job_Role": q(con, "SELECT * FROM dim_job_role"),
        "Dim_Grade": q(con, "SELECT * FROM dim_grade ORDER BY 1"),
        "Dim_Tenure_Band": q(con, "SELECT * FROM dim_tenure_band"),
    }
    sheets = {
        "Data_Quality": dq, "KPI_Summary": kpi, "Employee_Master": master, **dims,
        "Attrition_Drivers": pd.read_csv(PROC / "attrition_diagnostics.csv"),
        "Attrition_Cohort": q(con, "SELECT * FROM vw_attrition_cohort"),
        "Risk_Model_Drivers": pd.read_csv(PROC / "risk_drivers.csv"),
        "Early_Warning_Flags": pd.read_csv(PROC / "early_warning_flag_performance.csv"),
        "Critical_Role_Risk": pd.read_csv(PROC / "critical_role_risk.csv"),
        "Talent_Health": pd.read_csv(PROC / "talent_health_scorecard.csv"),
        "Headcount_Forecast": pd.read_csv(PROC / "headcount_forecast.csv"),
        "SYN_Attendance": q(con, "SELECT * FROM fct_attendance_monthly ORDER BY employee_id, month_start"),
        "SYN_Productivity": q(con, "SELECT * FROM syn_productivity_quarterly"),
        "SYN_Requisitions": q(con, "SELECT * FROM syn_requisitions"),
        "SYN_Hiring_By_Source": pd.read_csv(PROC / "hiring_by_source_syn.csv"),
        "SYN_Headcount_Monthly": q(con, """SELECT h.month_start, d.department, h.opening_headcount, h.hires, h.exits, h.is_synthetic
                                         FROM fct_headcount_monthly h JOIN dim_department d USING (department_key) ORDER BY 1, 2"""),
    }
    readme = pd.DataFrame({"PeoplePulse – workforce analytics workbook": [
        f"Snapshot date: {SNAPSHOT_DATE}. One table per sheet, header in row 1 – load directly into Power BI / Tableau / Looker Studio.",
        "REAL: Employee_Master and every sheet without the SYN_ prefix are computed from the 1,470 IBM HR Analytics records.",
        "SYNTHETIC: SYN_* sheets are generated (fixed seed) because the IBM dataset has no attendance, hiring or productivity history. Use them to demonstrate the method only.",
        "Employee_Master is unchanged from the source except for 4 appended model columns: risk_probability, risk_tier, top_risk_drivers, early_warning_flag_count.",
        "risk_probability = out-of-fold 12-month exit probability from a logistic-regression model (5-fold CV ROC AUC in README).",
        "Data_Quality lists issues found in the source workbook, including the KPI_Summary COUNTA bug (fixed here with live formulas).",
    ]})
    path = OUT / "PeoplePulse_Analytics.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="README", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        wb = xw.book
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            synthetic = ws.title.startswith("SYN_")
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="B45309" if synthetic else "1F3A5F")
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60 if ws.title == "README" else 40)
            if ws.title not in ("README",) and ws.max_row > 1:
                tbl = Table(displayName=f"t_{ws.title}", ref=ws.dimensions)
                tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
                ws.add_table(tbl)
        wb["README"].column_dimensions["A"].width = 160
        kws = wb["KPI_Summary"]
        for r in range(2, 7):
            kws[f"B{r}"].number_format = "0.0%" if r in (5, 6) else "#,##0"
        kws["B7"].number_format = "0.0%"
    print("wrote", path)


def charts(con):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "axes.spines.top": False, "axes.spines.right": False})
    def save(fig, name):
        fig.tight_layout(); fig.savefig(CHARTS / f"{name}.png", dpi=200); plt.close(fig)

    d = pd.read_csv(PROC / "attrition_diagnostics.csv")
    overall = json.loads((PROC / "model_metrics.json").read_text())["overall_exit_proportion"]
    sel = d[d.dimension.isin(["Job role", "Tenure band", "Overtime", "Manager tenure", "Grade", "Age band", "Job satisfaction"])]
    top = sel[sel.significant & (sel.exit_proportion > overall)].sort_values("exit_proportion", ascending=True).tail(10)
    fig, ax = plt.subplots(figsize=(8, 4.6))
    labels = top.dimension + ": " + top.value.astype(str)
    ax.barh(labels, top.exit_proportion * 100, color=CORAL, xerr=[(top.exit_proportion - top.ci_low) * 100, (top.ci_high - top.exit_proportion) * 100],
            error_kw={"ecolor": GREY, "capsize": 3})
    ax.axvline(overall * 100, color=NAVY, ls="--"); ax.text(overall * 100 + 0.5, len(top) - 0.35, f"Company {overall:.1%}", color=NAVY, fontsize=9)
    ax.set_xlabel("Exit proportion (%) with 95% CI"); ax.set_title("Cohorts with exit rates significantly above company level", loc="left", fontsize=11)
    save(fig, "attrition_hotspots")

    c = q(con, "SELECT * FROM vw_attrition_cohort")
    piv = c.pivot(index="department", columns="tenure_band", values="exit_proportion")[["0–1 Years", "2–3 Years", "4–5 Years", "6–10 Years", "10+ Years"]]
    cnt = c.pivot(index="department", columns="tenure_band", values="records")[piv.columns]
    fig, ax = plt.subplots(figsize=(8, 3.2))
    im = ax.imshow(piv.values * 100, cmap="OrRd", vmin=0, vmax=60, aspect="auto")
    ax.set_xticks(range(5), piv.columns); ax.set_yticks(range(len(piv)), piv.index)
    for i in range(piv.shape[0]):
        for j in range(piv.shape[1]):
            ax.text(j, i, f"{piv.values[i, j]:.0%}\n(n={int(cnt.values[i, j])})", ha="center", va="center", fontsize=9, color="white" if piv.values[i, j] > 0.45 else "black")
    fig.colorbar(im, ax=ax, label="Exit %"); ax.set_title("Exit proportion by department × tenure", loc="left")
    save(fig, "cohort_heatmap")

    f = pd.read_csv(PROC / "early_warning_flag_performance.csv")
    fig, ax = plt.subplots(figsize=(6, 3.6))
    ax.bar(f.flag_bucket.astype(str), f.exit_proportion * 100, color=[TEAL, TEAL, AMBER, CORAL, CORAL])
    for i, r in f.iterrows():
        ax.text(i, r.exit_proportion * 100 + 1, f"{r.exit_proportion:.0%}\n n={r.records}", ha="center", fontsize=9)
    ax.set_xlabel("Number of early-warning flags"); ax.set_ylabel("Exit proportion (%)"); ax.set_ylim(0, 60)
    ax.set_title("Exit rate rises steeply with flag count", loc="left")
    save(fig, "early_warning_flags")

    dr = pd.read_csv(PROC / "risk_drivers.csv").head(10).iloc[::-1]
    fig, ax = plt.subplots(figsize=(7, 4.2))
    ax.barh(dr.feature, dr.odds_ratio, color=[CORAL if x > 1 else TEAL for x in dr.odds_ratio])
    ax.axvline(1, color=NAVY, lw=1); ax.set_xlabel("Odds ratio (numeric features per 1 SD)")
    ax.set_title("Top attrition-risk model drivers", loc="left")
    save(fig, "risk_drivers")

    cr = pd.read_csv(PROC / "critical_role_risk.csv").sort_values("expected_exits_12m")
    fig, ax = plt.subplots(figsize=(7.5, 4.2))
    ax.barh(cr.job_role, cr.expected_exits_12m, color=[CORAL if x else GREY for x in cr.is_critical])
    for i, r in enumerate(cr.itertuples()):
        ax.text(r.expected_exits_12m + 0.5, i, f"{r.expected_exits_12m:.0f} of {r.active:.0f}", va="center", fontsize=9)
    ax.set_xlabel("Expected exits next 12 months (sum of risk probabilities)")
    ax.set_title("Expected exits by role (coral = critical role)", loc="left")
    save(fig, "critical_roles")

    fc = pd.read_csv(PROC / "headcount_forecast.csv").query("department == 'All'")
    fig, ax = plt.subplots(figsize=(8, 4))
    colors = [NAVY, CORAL, TEAL, AMBER]
    for (name, g), col in zip(fc.groupby("scenario", sort=False), colors):
        x = ["2026Q3 (actual)"] + list(g.quarter)
        y = [g.start_headcount.iloc[0]] + list(g.headcount_p50)
        ax.plot(x, y, marker="o", color=col, label=name)
        ax.fill_between(x[1:], g.headcount_p10, g.headcount_p90, color=col, alpha=0.12)
    ax.set_ylabel("Active headcount"); ax.legend(fontsize=8, frameon=False, loc="lower left")
    ax.set_title("Quarterly headcount forecast (median, P10–P90 band)", loc="left")
    save(fig, "headcount_forecast")

    h = pd.read_csv(PROC / "hiring_by_source_syn.csv")
    fig, ax = plt.subplots(figsize=(6.5, 3.6))
    ax.bar(h.hiring_source, h.avg_time_to_fill, color=TEAL)
    ax.set_ylabel("Avg time to fill (days)"); ax.set_title("SYNTHETIC: time to fill by hiring source", loc="left")
    plt.setp(ax.get_xticklabels(), rotation=15)
    save(fig, "hiring_source_syn")

    a = q(con, "SELECT month_start, sum(absent_days)/sum(scheduled_days) AS r FROM fct_attendance_monthly GROUP BY 1 ORDER BY 1")
    fig, ax = plt.subplots(figsize=(6.5, 3.2))
    ax.plot(a.month_start.dt.strftime("%b %y"), a.r * 100, color=NAVY, marker="o")
    ax.set_ylabel("Absenteeism %"); ax.set_ylim(0, 5); ax.set_title("SYNTHETIC: monthly absenteeism", loc="left")
    plt.setp(ax.get_xticklabels(), rotation=45)
    save(fig, "absenteeism_syn")
    print("charts ->", CHARTS)


def dashboard_json(con):
    def recs(df):
        return json.loads(df.to_json(orient="records", date_format="iso"))
    emp = q(con, """SELECT v.employee_id AS id, v.department_short AS dept, v.job_role AS role, v.role_family AS fam, v.grade, v.tenure_band AS tenure,
                           v.is_exit AS exit, CASE WHEN v.overtime THEN 1 ELSE 0 END AS ot, v.job_satisfaction AS sat,
                           v.years_with_current_manager = 0 AS newmgr, v.years_with_current_manager AS ywm, v.age_band AS age,
                           v.monthly_income AS income, r.risk_probability AS risk, r.risk_tier AS tier, r.top_risk_drivers AS drivers,
                           f.flag_count AS flags, round(a.absent / a.sched, 4) AS abs
                    FROM vw_employee v JOIN an_risk_scores r USING (employee_id) JOIN an_early_warning_flags f USING (employee_id)
                    JOIN (SELECT employee_id, sum(absent_days) AS absent, sum(scheduled_days) AS sched FROM fct_attendance_monthly GROUP BY 1) a
                      USING (employee_id)""")
    tables = q(con, """SELECT table_name AS name, table_type AS type,
                              (SELECT count(*) FROM information_schema.columns c WHERE c.table_name = t.table_name) AS cols
                       FROM information_schema.tables t WHERE table_name SIMILAR TO '(dim|fct|vw)_.*' ORDER BY 1""")
    tables["rows"] = [con.execute(f'SELECT count(*) FROM "{t}"').fetchone()[0] for t in tables.name]
    tables["synthetic"] = tables.name.isin(["fct_attendance_monthly", "fct_productivity_quarterly", "fct_requisition", "fct_headcount_monthly",
                                            "vw_absenteeism_monthly", "vw_hiring_by_source", "vw_productivity_by_role"])
    data = {
        "snapshot": SNAPSHOT_DATE,
        "employees": recs(emp),
        "metrics": json.loads((PROC / "model_metrics.json").read_text()),
        "drivers": recs(pd.read_csv(PROC / "risk_drivers.csv").head(12)),
        "flags": recs(pd.read_csv(PROC / "early_warning_flag_performance.csv")),
        "critical": recs(pd.read_csv(PROC / "critical_role_risk.csv")),
        "scorecard": recs(pd.read_csv(PROC / "talent_health_scorecard.csv")),
        "forecast": recs(pd.read_csv(PROC / "headcount_forecast.csv")),
        "absenteeism": recs(q(con, """SELECT strftime(month_start, '%Y-%m') AS month, d.department_short AS dept,
                                      round(100*sum(absent_days)/sum(scheduled_days), 2) AS rate
                                      FROM fct_attendance_monthly a JOIN dim_employee de USING (employee_id) JOIN dim_department d USING (department_key)
                                      GROUP BY ALL ORDER BY 1""")),
        "headcount": recs(q(con, """SELECT strftime(month_start, '%Y-%m') AS month, sum(opening_headcount) AS headcount, sum(hires) AS hires, sum(exits) AS exits
                                    FROM fct_headcount_monthly GROUP BY 1 ORDER BY 1""")),
        "hiring": recs(pd.read_csv(PROC / "hiring_by_source_syn.csv")),
        "productivity": recs(q(con, """SELECT review_period AS q, role_family AS fam, round(avg(avg_productivity), 1) AS score
                                       FROM vw_productivity_by_role GROUP BY ALL ORDER BY 1, 2""")),
        "prodByRating": recs(pd.read_csv(PROC / "productivity_by_rating_syn.csv")),
        "dq": json.loads((STAGE / "data_quality.json").read_text()),
        "kpi": recs(q(con, "SELECT * FROM vw_kpi_overall"))[0],
        "tables": recs(tables),
        "headcountDept": recs(q(con, """SELECT strftime(month_start, '%Y-%m') AS month, d.department_short AS dept, opening_headcount AS headcount, hires, exits
                                        FROM fct_headcount_monthly h JOIN dim_department d USING (department_key) ORDER BY 1, 2""")),
        "prodRole": recs(pd.read_csv(PROC / "productivity_by_role_syn.csv")),
        "ttfGrade": recs(pd.read_csv(PROC / "time_to_fill_by_grade_syn.csv")),
        "probation": recs(pd.read_csv(PROC / "probation_outcomes_syn.csv")),
        "firstYear": recs(pd.read_csv(PROC / "first_year_exits.csv")),
        "leaveExit": recs(pd.read_csv(PROC / "leave_vs_exit_syn.csv")),
        "workforceMeta": json.loads((PROC / "workforce_meta.json").read_text()),
    }
    path = ROOT / "dashboard" / "dashboard_data.json"
    path.parent.mkdir(exist_ok=True)
    path.write_text(json.dumps(data))
    print("wrote", path, f"{path.stat().st_size / 1024:.0f} KB")


def main():
    con = duckdb.connect(str(DB), read_only=True)
    workbook(con)
    charts(con)
    dashboard_json(con)
    con.close()


if __name__ == "__main__":
    main()
